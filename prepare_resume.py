"""把简历 docx 提取成纯文本，供 RAG 使用。"""
import sys
from docx import Document

src = sys.argv[1]
dst = sys.argv[2]

doc = Document(src)
lines = []
for p in doc.paragraphs:
    t = p.text.strip()
    if t:
        lines.append(t)
for t in doc.tables:
    for r in t.rows:
        cells = [c.text.strip() for c in r.cells]
        if any(cells):
            lines.append(' | '.join(cells))

text = '\n'.join(lines)
with open(dst, 'w', encoding='utf-8') as f:
    f.write(text)

print(f'简历文本已写入: {dst}')
print(f'  行数: {len(lines)}  字符数: {len(text)}')
print('--- 前 12 行 ---')
for l in lines[:12]:
    print('  ' + l[:90])
